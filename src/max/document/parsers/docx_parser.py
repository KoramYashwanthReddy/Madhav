"""DOCX document parser for Module 24 — Document Intelligence."""

import hashlib
import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, datetime

from max.document.domain.enums import DocumentFormat, DocumentSource, ElementType
from max.document.domain.models import (
    DocumentElement,
    DocumentExtractionResult,
    DocumentHeading,
    DocumentLocation,
    DocumentMetadata,
    DocumentPage,
    DocumentParagraph,
    DocumentProvenance,
    DocumentSection,
    DocumentStructure,
    DocumentTable,
    DocumentTableCell,
    DocumentTableRow,
)
from max.document.parsers.base import BaseDocumentParser

try:
    import docx
except ImportError:
    docx = None  # type: ignore[assignment]


class DocxDocumentParser(BaseDocumentParser):
    """Parses Word (.docx) documents into structured headings, paragraphs, and tables."""

    @property
    def supported_formats(self) -> list[DocumentFormat]:
        return [DocumentFormat.DOCX]

    def parse(
        self,
        document_id: str,
        version_id: str,
        content_bytes: bytes,
        filename: str,
        source_reference: str = "",
    ) -> DocumentExtractionResult:
        content_hash = hashlib.sha256(content_bytes).hexdigest()
        headings: list[DocumentHeading] = []
        sections: list[DocumentSection] = []
        paragraphs: list[DocumentParagraph] = []
        tables: list[DocumentTable] = []
        elements: list[DocumentElement] = []
        full_text_parts: list[str] = []
        author = title = None

        if docx is not None:
            # 1. Try python-docx
            try:
                doc = docx.Document(io.BytesIO(content_bytes))
                core_props = doc.core_properties
                author = core_props.author
                title = core_props.title

                for p_idx, p in enumerate(doc.paragraphs, start=1):
                    p_text = p.text.strip()
                    if not p_text:
                        continue
                    full_text_parts.append(p_text)
                    loc = DocumentLocation(page_number=1, row_index=p_idx)

                    style_name = getattr(p.style, "name", "") or ""
                    if style_name.startswith("Heading"):
                        try:
                            level = int(style_name.replace("Heading", "").strip())
                        except ValueError:
                            level = 1
                        hdg = DocumentHeading(text=p_text, level=level, location=loc)
                        headings.append(hdg)
                        sec = DocumentSection(title=p_text, level=level, content="")
                        sections.append(sec)

                        elem = DocumentElement(
                            document_id=document_id,
                            version_id=version_id,
                            element_type=ElementType.HEADING,
                            content=p_text,
                            location=loc,
                            metadata={"level": level},
                        )
                        elements.append(elem)
                    else:
                        para = DocumentParagraph(text=p_text, location=loc)
                        paragraphs.append(para)

                        elem = DocumentElement(
                            document_id=document_id,
                            version_id=version_id,
                            element_type=ElementType.PARAGRAPH,
                            content=p_text,
                            location=loc,
                        )
                        elements.append(elem)

                for t_idx, t in enumerate(doc.tables, start=1):
                    t_rows: list[DocumentTableRow] = []
                    headers: list[str] = []
                    for r_idx, row in enumerate(t.rows):
                        r_cells: list[DocumentTableCell] = []
                        for c_idx, cell in enumerate(row.cells):
                            cell_text = cell.text.strip()
                            r_cells.append(
                                DocumentTableCell(
                                    row_index=r_idx, column_index=c_idx, content=cell_text
                                )
                            )
                        if r_idx == 0:
                            headers = [c.content for c in r_cells]
                        else:
                            t_rows.append(DocumentTableRow(row_index=r_idx - 1, cells=r_cells))

                    tbl = DocumentTable(
                        headers=headers,
                        rows=t_rows,
                        row_count=len(t_rows),
                        column_count=len(headers),
                        location=DocumentLocation(page_number=1, row_index=t_idx),
                    )
                    tables.append(tbl)
                    elem = DocumentElement(
                        document_id=document_id,
                        version_id=version_id,
                        element_type=ElementType.TABLE,
                        content=f"Table ({tbl.row_count}x{tbl.column_count})",
                        location=DocumentLocation(page_number=1, row_index=t_idx),
                    )
                    elements.append(elem)
            except Exception:
                pass

        # 2. Zip XML parsing fallback if python-docx not installed or failed
        if not full_text_parts:
            try:
                with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
                    xml_content = z.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for elem_p in root.findall(".//w:p", ns):
                        texts = [t.text for t in elem_p.findall(".//w:t", ns) if t.text]
                        p_text = "".join(texts).strip()
                        if p_text:
                            full_text_parts.append(p_text)
                            loc = DocumentLocation(page_number=1)
                            para = DocumentParagraph(text=p_text, location=loc)
                            paragraphs.append(para)
                            elem = DocumentElement(
                                document_id=document_id,
                                version_id=version_id,
                                element_type=ElementType.PARAGRAPH,
                                content=p_text,
                                location=loc,
                            )
                            elements.append(elem)
            except Exception:
                pass

        raw_text = "\n\n".join(full_text_parts)
        words = raw_text.split()

        metadata = DocumentMetadata(
            filename=filename,
            normalized_filename=filename.lower().strip(),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size_bytes=len(content_bytes),
            content_hash=content_hash,
            author=author,
            title=title or filename,
            page_count=1,
            word_count=len(words),
            character_count=len(raw_text),
        )

        page = DocumentPage(page_number=1, text=raw_text, element_ids=[e.element_id for e in elements])
        structure = DocumentStructure(
            pages=[page],
            sections=sections,
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            elements=elements,
        )

        provenance = DocumentProvenance(
            document_id=document_id,
            version_id=version_id,
            source_type=DocumentSource.LOCAL_FILE,
            source_reference=source_reference or filename,
            content_hash=content_hash,
            processed_at=datetime.now(UTC),
            parser_name="DocxDocumentParser",
        )

        return DocumentExtractionResult(
            document_id=document_id,
            metadata=metadata,
            structure=structure,
            raw_text=raw_text,
            provenance=provenance,
        )
