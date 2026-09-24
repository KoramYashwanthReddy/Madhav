"""PPTX document parser for Module 24 — Document Intelligence."""

import hashlib
import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, datetime

from max.document.domain.enums import DocumentFormat, DocumentSource, ElementType
from max.document.domain.models import (
    DocumentElement,
    DocumentExtractionResult,
    DocumentLocation,
    DocumentMetadata,
    DocumentPage,
    DocumentProvenance,
    DocumentSlide,
    DocumentStructure,
)
from max.document.parsers.base import BaseDocumentParser

try:
    import pptx
except ImportError:
    pptx = None


class PptxDocumentParser(BaseDocumentParser):
    """Parses PowerPoint (.pptx) presentations into slides, titles, text blocks, and notes."""

    @property
    def supported_formats(self) -> list[DocumentFormat]:
        return [DocumentFormat.PPTX]

    def parse(
        self,
        document_id: str,
        version_id: str,
        content_bytes: bytes,
        filename: str,
        source_reference: str = "",
    ) -> DocumentExtractionResult:
        content_hash = hashlib.sha256(content_bytes).hexdigest()
        slides: list[DocumentSlide] = []
        elements: list[DocumentElement] = []
        full_text_parts: list[str] = []

        if pptx is not None:
            # 1. Try python-pptx
            try:
                prs = pptx.Presentation(io.BytesIO(content_bytes))
                for s_idx, slide in enumerate(prs.slides, start=1):
                    slide_title = ""
                    text_blocks: list[str] = []
                    notes_text = ""

                    if slide.notes_slide and slide.notes_slide.notes_text_frame:
                        notes_text = slide.notes_slide.notes_text_frame.text.strip()

                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            stext = shape.text.strip()
                            text_blocks.append(stext)
                            if shape == slide.shapes.title:
                                slide_title = stext

                    slide_text = "\n".join(text_blocks)
                    full_text_parts.append(f"--- Slide {s_idx}: {slide_title} ---\n{slide_text}")

                    slide_obj = DocumentSlide(
                        slide_number=s_idx,
                        title=slide_title,
                        text_blocks=text_blocks,
                        speaker_notes=notes_text,
                    )
                    slides.append(slide_obj)

                    elem = DocumentElement(
                        document_id=document_id,
                        version_id=version_id,
                        element_type=ElementType.SLIDE,
                        content=f"Slide {s_idx}: {slide_title or 'Untitled'}",
                        location=DocumentLocation(page_number=1, slide_number=s_idx),
                        metadata={"title": slide_title, "blocks": len(text_blocks)},
                    )
                    elements.append(elem)
            except Exception:
                pass

        # 2. Zip XML fallback
        if not slides:
            try:
                with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
                    slide_files = sorted(
                        [f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")]
                    )
                    for idx, sfile in enumerate(slide_files, start=1):
                        s_data = z.read(sfile)
                        s_root = ET.fromstring(s_data)
                        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                        texts = [t.text for t in s_root.findall(".//a:t", ns) if t.text]
                        s_text = " ".join(texts)

                        full_text_parts.append(f"--- Slide {idx} ---\n{s_text}")
                        slide_obj = DocumentSlide(
                            slide_number=idx,
                            title=f"Slide {idx}",
                            text_blocks=texts,
                        )
                        slides.append(slide_obj)

                        elem = DocumentElement(
                            document_id=document_id,
                            version_id=version_id,
                            element_type=ElementType.SLIDE,
                            content=f"Slide {idx}",
                            location=DocumentLocation(page_number=1, slide_number=idx),
                        )
                        elements.append(elem)
            except Exception:
                pass

        raw_text = "\n\n".join(full_text_parts)
        words = raw_text.split()

        metadata = DocumentMetadata(
            filename=filename,
            normalized_filename=filename.lower().strip(),
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            file_size_bytes=len(content_bytes),
            content_hash=content_hash,
            title=filename,
            page_count=1,
            slide_count=len(slides),
            word_count=len(words),
            character_count=len(raw_text),
        )

        page = DocumentPage(page_number=1, text=raw_text, element_ids=[e.element_id for e in elements])
        structure = DocumentStructure(
            pages=[page],
            slides=slides,
            elements=elements,
        )

        provenance = DocumentProvenance(
            document_id=document_id,
            version_id=version_id,
            source_type=DocumentSource.LOCAL_FILE,
            source_reference=source_reference or filename,
            content_hash=content_hash,
            processed_at=datetime.now(UTC),
            parser_name="PptxDocumentParser",
        )

        return DocumentExtractionResult(
            document_id=document_id,
            metadata=metadata,
            structure=structure,
            raw_text=raw_text,
            provenance=provenance,
        )
